import os
import shutil
from pathlib import Path
import pandas as pd
import logging
from openpyxl import Workbook

logging.basicConfig(
    level = logging.INFO,
    format = "%(asctime)s - %(levelname)s - %(message)s",
    filename = "Sales_Automation.log",
    force = True
)
logging.info("Automation Started")

input_folder = Path("Incoming_Data")
csv_file = list(input_folder.glob("*.csv"))
logging.info("CSV file loaded")

df = pd.read_csv(csv_file[0])
logging.info("pandas DataFrame Created")

df.columns = df.columns.str.lower().str.replace(" ", "_")
df = df.sort_values(by = "date")
logging.info("Columns name standerdized and data has been sorted")


df["unit_price"] = df.groupby("product")["unit_price"].transform(lambda x: x.fillna(x.mean()))
df["quantity"] = df.groupby("product")["quantity"].transform(lambda x: x.fillna(x.median()))
logging.info("Missing values filled")

df.drop_duplicates(inplace=True)
logging.info("Duplicate values removed")


df["sales"] = df["quantity"]*df["unit_price"]

df["date"] = pd.to_datetime(df["date"])
df["day_name"] = df["date"].dt.day_name()
df["month_name"] = df["date"].dt.month_name()
logging.info("New columns created")


products_summary= df.groupby("product").agg({"sales":"sum", "quantity":"count"}).reset_index()
regions_summary = df.groupby("region")["sales"].sum().reset_index()
days_summary = df.groupby("day_name")["sales"].sum().reset_index()
months_summary = df.groupby("month_name")["sales"].sum().reset_index()
logging.info("Summary data created")


wb = Workbook()
ws_data = wb.active
ws_data.title = "raw_data"
ws_products = wb.create_sheet("products_summary")
ws_regions = wb.create_sheet("regions_summary")
ws_months = wb.create_sheet("months_summary")
ws_days = wb.create_sheet("days_summary")
logging.info("Workbook created with all worksheets")


from openpyxl.utils.dataframe import dataframe_to_rows

for i in dataframe_to_rows(df, index = False, header = True):
    ws_data.append(i)

for j in dataframe_to_rows(products_summary, index = False, header = True):
    ws_products.append(j)

for k in dataframe_to_rows(regions_summary, index = False, header = True):
    ws_regions.append(k)

for l in dataframe_to_rows(months_summary, index = False, header = True):
    ws_months.append(l)

for m in dataframe_to_rows(days_summary, index = False, header = True):
    ws_days.append(m)

logging.info("Dataframes has been processed to Excel")

from openpyxl.styles import Font
from openpyxl.styles import PatternFill

def excel_work(ws):

    bold = Font(bold = True, color = "FFFFFF")
    header_color = PatternFill(fgColor = "1F4E78", fill_type = "solid")


    for cell in ws[1]:
        cell.font = bold
        cell.fill = header_color
    logging.info("Headers converted to bold and color has been changed")

    ws.auto_filter.ref = ws.dimensions
    logging.info("Filter buttons activated")


    for col in ws.columns:
        l = []
        for cell in col:
            if cell.value:
                val_len = len(str(cell.value))
            else:
                val_len = 0
            l.append(val_len)
        length = max(l)
        
        ws.column_dimensions[col[0].column_letter].width = length+2
        
    logging.info("Column width auto adjusted")

    ws.freeze_panes = "A2"
    logging.info("First row freezed")
    return ws

from openpyxl.formatting.rule import ColorScaleRule, DataBarRule

def databar(ws):
    last_row = ws.max_row
    data_rule = DataBarRule(start_type = "min", end_type = "max", color = "638EC6", showValue = True)
    
    ws.conditional_formatting.add(f"B2:B{last_row}", data_rule)
    return ws

from openpyxl.chart import BarChart, LineChart, Reference

def bar_chart(ws):
    last_row = ws.max_row
    
    data = Reference(ws, min_col = 2, min_row = 1, max_row = last_row, max_col = 2)
    cats = Reference(ws, min_col = 1, min_row = 1, max_row = last_row)
    
    bar = BarChart()
    bar.title = "Total Revenue Distribution"
    bar.y_axis.title = "Total Revenue"
    
    bar.add_data(data = data, titles_from_data = True)
    bar.set_categories(cats)
    
    ws.add_chart(bar, "G2")
    return ws


def line_chart(ws):
    last_row = ws.max_row
    
    data = Reference(ws, min_col = 2, min_row = 1, max_row = last_row, max_col = 2)
    cats = Reference(ws, min_col = 1, min_row = 1, max_row = last_row)
    
    line = LineChart()
    line.title = "Total Revenue Distribution"
    line.y_axis.title = "Total Revenue"
    
    line.add_data(data = data, titles_from_data = True)
    line.set_categories(cats)
    
    ws.add_chart(line, "G2")
    return ws



excel_work(ws_data)
excel_work(ws_products)
excel_work(ws_regions)
excel_work(ws_months)
excel_work(ws_days)
logging.info("Excel formatting done on all sheets")

databar(ws_products)
databar(ws_regions)
databar(ws_months)
databar(ws_days)
logging.info("data bar added to the B column")

bar_chart(ws_products)
bar_chart(ws_months)
line_chart(ws_regions)
line_chart(ws_days)
logging.info("Charts has been added to the summary sheets")


os.makedirs("Output", exist_ok = True)
wb.save("Output/Sales_Automation.xlsx")
logging.info("File has been saved to the Output folder")


os.makedirs("Archive_Data", exist_ok = True)
shutil.move(csv_file[0], "Archive_Data")
logging.info("Raw data has been moved to Archive folder")



import yagmail

def send_mail():
    sender = "ADD YOUR EMAIL HERE"
    receiver = "ADD TARGET EMAIL HERE"
    password = "ADD YOUR PASSWORD HERE" #App password
    file_path = "Output/Sales_Automation.xlsx"
    
    yag = yagmail.SMTP(user = sender, password = password)
    yag.send(to = receiver,
             subject = "Sales Report",
             contents = "Automated Sales Report for Today",
             attachments = file_path)
    
    logging.info("Final report sent to the Receiver")
    
send_mail()
logging.info("Automation Completed")
